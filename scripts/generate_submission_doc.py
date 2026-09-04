import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Color Palette
NAVY = RGBColor(11, 25, 44)       # #0B192C
BLUE = RGBColor(58, 134, 255)     # #3A86FF
PURPLE = RGBColor(108, 99, 255)   # #6C63FF
SLATE = RGBColor(100, 116, 139)   # #64748B
DARK = RGBColor(30, 41, 59)       # #1E293B

# Helper: set cell background shading
def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Helper: set cell margins/padding
def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# ==========================================
# HEADER & TITLE
# ==========================================
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_top.paragraph_format.space_after = Pt(2)

r_sub = p_top.add_run('AMITY UNIVERSITY MUMBAI • AMI-ASTRO ASTRONOMY CLUB\n')
r_sub.font.size = Pt(10)
r_sub.font.bold = True
r_sub.font.color.rgb = PURPLE

r_title = p_top.add_run('Website Management Team — Selection Task Submission')
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = NAVY

p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(12)
r_author = p_meta.add_run('Candidate / Developer: ')
r_author.bold = True
p_meta.add_run('Rajdeep  |  ')
r_role = p_meta.add_run('Role: ')
r_role.bold = True
p_meta.add_run('Website Management Team  |  ')
r_date = p_meta.add_run('Date: ')
r_date.bold = True
p_meta.add_run('September 2026')

# ==========================================
# CALLOUT BOX: LIVE URL & GITHUB
# ==========================================
callout_table = doc.add_table(rows=1, cols=1)
callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = callout_table.cell(0, 0)
set_cell_shading(cell, 'F8FAFC')
set_cell_margins(cell, top=180, bottom=180, left=220, right=220)

p_callout = cell.paragraphs[0]
p_callout.paragraph_format.space_before = Pt(4)
p_callout.paragraph_format.space_after = Pt(4)

r_url_label = p_callout.add_run('🌐 Live Deployed Functional Website:\n')
r_url_label.font.size = Pt(11)
r_url_label.font.bold = True
r_url_label.font.color.rgb = NAVY

r_url = p_callout.add_run('https://ami-astro-amber.vercel.app/\n')
r_url.font.size = Pt(14)
r_url.font.bold = True
r_url.font.color.rgb = BLUE

r_repo_label = p_callout.add_run('📁 GitHub Repository: ')
r_repo_label.font.bold = True
r_repo = p_callout.add_run('https://github.com/rajdeep-r24/ami-astro\n')
r_repo.font.color.rgb = PURPLE

r_status = p_callout.add_run('🚀 Deployed on Vercel Global Edge Network with Integrated Real-Time Web Analytics')
r_status.font.size = Pt(9.5)
r_status.font.italic = True
r_status.font.color.rgb = SLATE

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ==========================================
# 1. EXECUTIVE SUMMARY
# ==========================================
h1 = doc.add_heading(level=1)
r_h1 = h1.add_run('1. Executive Summary')
r_h1.font.color.rgb = NAVY
r_h1.font.size = Pt(15)

p_exec = doc.add_paragraph(
    'In response to the selection requirements for the Website Management Team, a complete, state-of-the-art, '
    'and fully responsive web application was designed, engineered, and deployed for the AMI-ASTRO Astronomy Club '
    'of Amity University Mumbai. Rather than presenting a static wireframe or mock-up, this submission delivers an active, '
    'production-grade Next.js 16 web application featuring an interactive 3D solar system, dynamic astronomical event '
    'calendars, high-resolution astrophotography gallery archives, official announcement dispatches, and an integrated '
    'registration and inquiry system.'
)
p_exec.paragraph_format.line_spacing = 1.15
p_exec.paragraph_format.space_after = Pt(10)

# ==========================================
# 2. DELIVERABLES CHECKLIST TABLE
# ==========================================
h2 = doc.add_heading(level=1)
r_h2 = h2.add_run('2. Selection Task Deliverables & Compliance Audit')
r_h2.font.color.rgb = NAVY
r_h2.font.size = Pt(15)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Required Deliverable'
hdr_cells[1].text = 'Implementation in Functional Prototype'
hdr_cells[2].text = 'Status'

for c in hdr_cells:
    set_cell_shading(c, '0B192C')
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)

deliverables = [
    (
        'Home Page',
        'Interactive Three.js 3D solar system canvas, natural star twinkle, landscape silhouette, club overview, upcoming celestial events preview, latest announcements, and quick-action registration triggers.',
        '100% Complete'
    ),
    (
        'About Section / Page',
        'Dedicated /about page featuring the club chronicle timeline, mission and vision, telescope optics inventory (10" Dobsonians, refractors, solar filters), outreach programs, and interactive constellation signature.',
        '100% Complete'
    ),
    (
        'Teams or Committee Page',
        'Dedicated /team page showcasing the Faculty Mentors, Core Council, and Domain Leads with academic departments, bios, roles, and direct official Instagram links (@ami_astro_mumbai).',
        '100% Complete'
    ),
    (
        'Events & Registration Section',
        'Dedicated /events calendar page featuring the official September 2026 Astronomical Events Calendar (planetary conjunctions, Venus peak, equinox, harvest moon, zodiacal light) with category filters and official poster lightbox.',
        '100% Complete'
    ),
    (
        'Sample Event Page with Registration',
        '6 dynamic event routes (/events/[id]) featuring high-resolution composite banners, hour-by-hour schedules, prerequisites, contact leads, and an interactive digital pass registration modal.',
        '100% Complete'
    ),
    (
        'Gallery or Media Page',
        'Dedicated /gallery page featuring deep-sky astrophotography, field expeditions, planetary captures, category filtering tabs, optical telescope gear tags, and high-res lightbox modals.',
        '100% Complete'
    ),
    (
        'Announcements / Blog Section',
        'Dedicated /announcements page with real-time keyword search filter, category tags (Observation Alert, Workshop, News), and dynamic article reading routes (/announcements/[id]).',
        '100% Complete'
    ),
    (
        'Contact & Social Media Links',
        'Dedicated /contact page with compulsory Enrollment No validation, FAQ accordions, campus location map, and direct links to official @ami_astro_mumbai, YouTube, Discord, and Email.',
        '100% Complete'
    ),
    (
        'Mobile & Desktop Responsive',
        'Fully responsive across all device breakpoints (375px mobile, tablet, 1920px desktop) with mobile drawer navigation, touch gestures, and zero layout shift.',
        '100% Complete'
    ),
    (
        'Official Tab Branding (Favicon)',
        'Custom circular AMI-ASTRO emblem multi-resolution favicon (.ico, .png, apple-touch-icon) replacing default template icons across all browser tabs and mobile bookmarks.',
        '100% Complete'
    ),
]

for req, impl, status in deliverables:
    row_cells = table.add_row().cells
    row_cells[0].text = req
    row_cells[1].text = impl
    row_cells[2].text = status
    for i, c in enumerate(row_cells):
        set_cell_margins(c, top=80, bottom=80, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.1
        if i == 0:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)
        elif i == 1:
            p.runs[0].font.size = Pt(9)
        elif i == 2:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(16, 185, 129)
            p.runs[0].font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ==========================================
# 3. UI/UX & TECHNICAL ARCHITECTURE
# ==========================================
h3 = doc.add_heading(level=1)
r_h3 = h3.add_run('3. UI/UX Design Philosophy & Technical Architecture')
r_h3.font.color.rgb = NAVY
r_h3.font.size = Pt(15)

p_arch = doc.add_paragraph()
p_arch.paragraph_format.line_spacing = 1.15
p_arch.add_run('• Curated Celestial Aesthetics: ').bold = True
p_arch.add_run(
    'Built with a tailored dark-sky color system featuring obsidian (#070A14), electric nebula purple (#6C63FF), '
    'soft starlight violet (#A78BFA), and luminous cyan (#38BDF8) accents. Glassmorphic card surfaces provide depth '
    'and visual hierarchy without hindering text legibility.\n'
)
p_arch.add_run('• Interactive 3D WebGL Canvas: ').bold = True
p_arch.add_run(
    'The home hero features a custom Three.js astronomical canvas allowing users to inspect planetary orbits and solar corona '
    'glows, complemented by a subtle twinkling starfield and mountain landscape silhouette.\n'
)
p_arch.add_run('• Centralized Data Schemas: ').bold = True
p_arch.add_run(
    'All club content (events, announcements, members, equipment, gallery captures) is abstracted into modular TypeScript '
    'schemas (src/data/). UI components consume structured interfaces, allowing instant content updates without layout regressions.\n'
)
p_arch.add_run('• Strict Registration Pass Validation: ').bold = True
p_arch.add_run(
    'All event passes and membership registration forms enforce compulsory validation for Full Name, Email Address, '
    'Enrollment No, and Academic Department, ensuring accurate student records for campus observatory access.'
)

# ==========================================
# 4. ANNUAL WEBSITE MAINTENANCE PLAN
# ==========================================
h4 = doc.add_heading(level=1)
r_h4 = h4.add_run('4. Plan for Regular Website Maintenance & Updates Throughout the Year')
r_h4.font.color.rgb = NAVY
r_h4.font.size = Pt(15)

p_plan_intro = doc.add_paragraph(
    'As part of the Website Management Team, I will ensure the AMI-ASTRO platform remains dynamic, up-to-date, '
    'and accurate through the following structured operational workflow:'
)
p_plan_intro.paragraph_format.line_spacing = 1.15

p_plan = doc.add_paragraph()
p_plan.paragraph_format.line_spacing = 1.15
p_plan.add_run('1. Content & Editorial Synchronization:\n').bold = True
p_plan.add_run(
    '   • Coordinate on a bi-weekly cycle with the Digital Media and Editorial teams to ingest raw telescope captures, '
    'post-processed astrophotography, and session reports as soon as events conclude.\n'
    '   • Publish the monthly astronomical calendar (conjunctions, meteor showers, moon phases) 5 days prior to the start of each month.\n\n'
)
p_plan.add_run('2. Streamlined Data Architecture:\n').bold = True
p_plan.add_run(
    '   • The website utilizes a centralized, modular TypeScript data schema. Updating event dates, registration links, team council '
    'rosters, or gallery photos requires updating structured data objects without touching UI styling, eliminating deployment risks '
    'and layout breakage.\n\n'
)
p_plan.add_run('3. Form & Link Health Monitoring:\n').bold = True
p_plan.add_run(
    '   • Conduct periodic audits of registration forms, social media links, and external references to ensure zero dead links, accurate '
    'venue details, and working student registration passes with compulsory Enrollment Numbers.\n\n'
)
p_plan.add_run('4. Continuous Integration & Staging:\n').bold = True
p_plan.add_run(
    '   • Follow a disciplined Git branch workflow (feature branches for upcoming events/campaigns, automated build checks on pull requests) '
    'to preview changes before publishing to production, guaranteeing 100% uptime and zero breaking changes.\n\n'
)
p_plan.add_run('5. Performance & Analytics Monitoring:\n').bold = True
p_plan.add_run(
    '   • Monitor real-time page views and visitor trends using the integrated @vercel/analytics dashboard to identify which events and gallery '
    'categories attract the most student engagement, optimizing future promotions accordingly.\n'
    '   • Maintain compressed WebP/optimized JPEG asset delivery so the website loads instantly even on cellular networks during field expeditions.'
)

# ==========================================
# 5. FUTURE BACKEND INTEGRATION ROADMAP
# ==========================================
h5 = doc.add_heading(level=1)
r_h5 = h5.add_run('5. Future Backend Integration Roadmap')
r_h5.font.color.rgb = NAVY
r_h5.font.size = Pt(15)

p_future = doc.add_paragraph(
    'The website has been architected specifically to support seamless transition from frontend prototype to full-stack platform:\n'
    '• API Route Handlers: Next.js App Router natively supports Server Actions and Route Handlers (src/app/api/events, src/app/api/register).\n'
    '• Database Integration: The modular TypeScript models map 1:1 to database tables in Supabase, Firebase, or PostgreSQL with Prisma.\n'
    '• Automated Digital Passes: Ready to connect automated QR code ticket generation and email dispatch for registered attendees.'
)
p_future.paragraph_format.line_spacing = 1.15

# ==========================================
# 6. INTELLECTUAL PROPERTY & COPYRIGHT
# ==========================================
h6 = doc.add_heading(level=1)
r_h6 = h6.add_run('6. Intellectual Property & Copyright Notice')
r_h6.font.color.rgb = NAVY
r_h6.font.size = Pt(15)

p_lic = doc.add_paragraph(
    'Copyright © 2026 Rajdeep (rajdeep-r24). All rights reserved.\n'
    'This project, source code, 3D graphic shaders, custom UI components, and documentation are proprietary work '
    'submitted exclusively for evaluation by the Amity University Mumbai Website Management Team selection committee. '
    'No part of this codebase may be copied, reproduced, or distributed without prior written permission.'
)
p_lic.paragraph_format.line_spacing = 1.15

output_path = r'c:\Users\User\Desktop\1Ami-Astro\AMI_ASTRO_Website_Management_Submission.docx'
doc.save(output_path)
print(f'Document successfully created at {output_path}')
